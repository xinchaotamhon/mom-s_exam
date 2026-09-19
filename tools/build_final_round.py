"""Extract the final-round Word sources into a stable, provenance-aware JSON file.

This is intentionally separate from the existing 538-question pipeline.  The two
source documents are copied unchanged to data/source and are never rewritten.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "source"
OUT = ROOT / "data" / "derived" / "final-round.json"
CORRECTIONS_PATH = ROOT / "data" / "curated" / "final-round-corrections.json"
MCQ_PATH = SOURCE_DIR / "final-30-trac-nghiem.docx"
SCENARIO_PATH = SOURCE_DIR / "final-20-tinh-huong.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SECTION_NAMES = {
    "I": ("van-phong", "Công tác văn phòng"),
    "II": ("tuyen-giao-dan-van", "Công tác tuyên giáo và dân vận"),
    "III": ("noi-chinh", "Công tác nội chính"),
    "IV": ("kiem-tra-giam-sat", "Công tác kiểm tra, giám sát, thi hành kỷ luật Đảng"),
    "V": ("to-chuc-xay-dung-dang", "Công tác tổ chức xây dựng Đảng"),
}


def clean(text: str) -> str:
    text = text.replace("\xa0", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text)
    # Page-number artifacts accidentally present in a few source paragraphs.
    text = re.sub(r"(?<=:)\s+\d{1,4}$", "", text)
    return text.strip()


def strip_layout_artifact(text: str) -> str:
    """Remove page numbers accidentally appended to known answer/heading lines.

    This is deliberately narrow: numeric content in ordinary prompts is preserved.
    """
    text = re.sub(r"^(Đáp án\s*:\s*.+?\.)\s+\d{1,3}$", r"\1", text, flags=re.I)
    text = re.sub(r"^(\d+\.?\s*Tình huống\s*\d+\s*:)\s+\d{1,3}$", r"\1", text, flags=re.I)
    return text


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def doc_meta(path: Path) -> dict:
    doc = Document(path)
    pages = None
    with zipfile.ZipFile(path) as archive:
        app = ET.fromstring(archive.read("docProps/app.xml"))
        pages_node = next((node for node in app.iter() if node.tag.rsplit("}", 1)[-1] == "Pages"), None)
        if pages_node is not None and pages_node.text and pages_node.text.isdigit():
            pages = int(pages_node.text)
    return {
        "path": path.as_posix(),
        "sha256": sha256(path),
        "bytes": path.stat().st_size,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "pagesInAppXml": pages,
    }


def section_for(text: str):
    match = re.match(r"^([IVX]+)\.\s+(.+?)\s*\(\d+\s*câu\)", text, re.I)
    if not match:
        return None
    roman = match.group(1).upper()
    return SECTION_NAMES.get(roman)


def extract_mcq(path: Path):
    lines = [strip_layout_artifact(clean(p.text)) for p in Document(path).paragraphs if clean(p.text)]
    questions = []
    section_id, section_label = None, None
    current = None
    for line in lines:
        section = section_for(line)
        if section:
            section_id, section_label = section
            continue
        match = re.match(r"^Câu\s*(\d+)\.\s*(.*)$", line, re.I)
        if match:
            if current:
                questions.append(current)
            number = int(match.group(1))
            current = {
                "id": f"final-mcq-{number}",
                "number": number,
                "sectionId": section_id,
                "section": section_label,
                "prompt": match.group(2).strip(),
                "options": [],
                "answerReference": None,
            }
            continue
        option_matches = list(re.finditer(r"(?:^|\s)([A-D])\.\s*(.*?)(?=\s+[A-D]\.\s+|$)", line))
        if option_matches and current:
            for option in option_matches:
                current["options"].append({"id": option.group(1), "text": option.group(2).strip()})
            continue
        answer = re.match(r"^Đáp án\s*:?\s*([A-D])\s*[,(]?\s*(.*)$", line, re.I)
        if answer and current:
            current["correctOption"] = answer.group(1).upper()
            current["answerReference"] = answer.group(2).strip(" ,).") or None
    if current:
        questions.append(current)
    for question in questions:
        question["explanation"] = (
            f"Đáp án {question['correctOption']}. "
            + next(opt["text"] for opt in question["options"] if opt["id"] == question["correctOption"])
        )
        question["memoryCue"] = "Đọc kỹ căn cứ trong câu hỏi → đối chiếu 4 phương án → chọn phương án đúng theo tài liệu."
        question["verification"] = {
            "status": "source-provided",
            "sourceIds": ["final-mcq-source"],
            "note": "Đáp án và căn cứ được chép theo tài liệu trắc nghiệm cuối vòng; chưa tự xác minh độc lập.",
        }
    return questions


def extract_scenarios(path: Path):
    lines = [strip_layout_artifact(clean(p.text)) for p in Document(path).paragraphs if clean(p.text)]
    scenarios = []
    section_id, section_label = None, None
    current = None
    answer_start = False
    for line in lines:
        section = section_for(line)
        if section:
            section_id, section_label = section
            continue
        match = re.match(r"^(?:(\d+)\.\s*)?Tình huống\s*(\d+)\s*:?(.*)$", line, re.I)
        if match:
            if current:
                scenarios.append(current)
            number = int(match.group(2))
            current = {
                "id": f"final-scenario-{number}",
                "number": number,
                "sectionId": section_id,
                "section": section_label,
                "promptParts": [],
                "answerParts": [],
                "answerKind": None,
            }
            if match.group(3).strip():
                current["promptParts"].append(match.group(3).strip())
            answer_start = False
            continue
        if current is None:
            continue
        is_answer = bool(re.match(r"^(Đáp án|Gợi ý xử lý|Gợi ý giải pháp):?", line, re.I))
        if is_answer:
            answer_start = True
            current["answerKind"] = "answer" if line.lower().startswith("đáp án") else "practical-guidance"
            current["answerParts"].append(re.sub(r"^(Đáp án|Gợi ý xử lý|Gợi ý giải pháp):?\s*", "", line, flags=re.I))
        elif answer_start:
            current["answerParts"].append(line)
        else:
            current["promptParts"].append(line)
    if current:
        scenarios.append(current)
    for scenario in scenarios:
        scenario["prompt"] = "\n\n".join(scenario.pop("promptParts"))
        scenario["answer"] = "\n\n".join(scenario.pop("answerParts"))
        scenario["memoryCue"] = "Kết luận trước → nêu căn cứ → việc cần làm ngay → báo cáo, theo dõi và phòng ngừa."
        scenario["verification"] = {
            "status": "source-provided",
            "sourceIds": ["final-scenario-source"],
            "note": "Đáp án hoặc gợi ý xử lý được chép theo tài liệu tình huống cuối vòng; chưa tự xác minh độc lập.",
        }
    return scenarios


def main():
    for path in (MCQ_PATH, SCENARIO_PATH):
        if not path.exists():
            raise SystemExit(f"Missing immutable source: {path}")
    mcq = extract_mcq(MCQ_PATH)
    scenarios = extract_scenarios(SCENARIO_PATH)
    corrections_payload = json.loads(CORRECTIONS_PATH.read_text(encoding="utf-8"))
    corrections = {item["questionId"]: item for item in corrections_payload["corrections"]}
    for question in mcq:
        correction = corrections.get(question["id"])
        if correction and correction.get("correctedOption"):
            question["sourceAnswer"] = question["correctOption"]
            question["correctOption"] = correction["correctedOption"]
            selected = next(option for option in question["options"] if option["id"] == question["correctOption"])
            question["explanation"] = f"Đáp án {question['correctOption']}. {selected['text']}"
            question["correction"] = {
                "note": correction["correctionNote"],
                "sourceIds": correction["sourceIds"],
                "status": "publicly-verified",
            }
            question["verification"] = {
                "status": "publicly-verified",
                "sourceIds": ["final-mcq-source", *correction["sourceIds"]],
                "note": (
                    "Đáp án đã đối chiếu với văn bản công khai chính thức; bản gốc trong tài liệu vẫn được lưu để đối chiếu."
                    if question["sourceAnswer"] == question["correctOption"]
                    else "Đáp án đã hiệu chỉnh theo văn bản công khai chính thức; bản gốc trong tài liệu vẫn được lưu để đối chiếu."
                ),
            }
    if len(mcq) != 30:
        raise SystemExit(f"Expected 30 final MCQ, found {len(mcq)}")
    if len(scenarios) != 20:
        raise SystemExit(f"Expected 20 final scenarios, found {len(scenarios)}")
    for q in mcq:
        if len(q["options"]) != 4 or q.get("correctOption") not in "ABCD":
            raise SystemExit(f"Invalid final MCQ: {q['id']}")
    if any(not q["prompt"] or not q["answer"] for q in scenarios):
        raise SystemExit("A final scenario is missing prompt or source answer/guidance")
    for scenario in scenarios:
        scenario["sourceAnswer"] = scenario["answer"]
        correction = corrections.get(scenario["id"])
        if correction:
            scenario["answer"] = correction["correctedAnswer"]
            scenario["correction"] = {
                "note": correction["correctionNote"],
                "sourceIds": correction["sourceIds"],
                "status": "publicly-verified",
            }
            scenario["verification"] = {
                "status": "publicly-verified",
                "sourceIds": ["final-scenario-source", *correction["sourceIds"]],
                "note": "Lời giải hiển thị đã hiệu chỉnh theo văn bản công khai chính thức; bản gốc trong tài liệu vẫn được lưu để đối chiếu.",
            }
    payload = {
        "schemaVersion": 1,
        "generatedAt": date.today().isoformat(),
        "provenance": {
            "method": "python-docx and OOXML structure extraction; source documents copied byte-for-byte",
            "renderStatus": "not-rendered: bundled LibreOffice unavailable on host",
            "sources": {
                "final-mcq-source": doc_meta(MCQ_PATH),
                "final-scenario-source": doc_meta(SCENARIO_PATH),
            },
            "correctionSourceIds": [source["id"] for source in corrections_payload["sources"]],
            "correctionsSha256": sha256(CORRECTIONS_PATH),
        },
        "stats": {"multipleChoice": len(mcq), "scenarios": len(scenarios), "total": len(mcq) + len(scenarios)},
        "multipleChoice": mcq,
        "scenarios": scenarios,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": OUT.as_posix(), "stats": payload["stats"], "sources": payload["provenance"]["sources"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
